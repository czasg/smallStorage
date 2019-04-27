import xlrd,xlwt
from docx import Document
"""
Fail
"""
excel = xlwt.Workbook()
excel_sheet = excel.add_sheet('table')

doc = Document('text.docx')
table = doc.tables[0]

nrows = len(table.column_cells(0))
for i in range(nrows):
    row = table.column_cells(i)
    texts = [r.text for r in row]
    for m,n in enumerate(texts):
        excel_sheet.write(i,m,n)

excel.save('cza.xls')
r = xlrd.open_workbook('cza.xls')
sheet = r.sheet_by_name("table")
print(sheet)